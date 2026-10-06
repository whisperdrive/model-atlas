"""Which rows of a model are its statements, found by the identities that hold in its saved values.

A financial model has a shape: a P&L chain (revenue + costs = EBITDA ... NPAT), a balance sheet that balances, a cash
flow whose closing cash is the balance sheet's cash, debt and fixed-asset corkscrews whose closing balances are balance
sheet lines, a sources-and-uses block that nets to zero, revenue built from segments, distributions that agree across the
DCF, the cash flow and the equity roll-forward. modelatlas/ontology.py lists those identities. This module finds the rows
that play the roles, and the identities are the detection signal, not an after-the-fact test:

  1. the model's own check rows are mined first (`checks`, `given`): a check formula names rows that must agree, and the
     check's saved value is the model's verdict. A binding that a check names wins every tie.
  2. structure is read from the R1C1 patterns (rows.patterns) and the live edges: subtotals (a row that sums same-column
     rows), corkscrews (a closing row = opening + movements where the opening reads the closing one column back),
     consolidations (a row that sums the same row of several sheets), links (a row that is one other row).
  3. roles are bound where an independent identity ties two structures by VALUE across every period: assets = liabilities
     + equity picks the balance sheet's three totals when something independent of the values says it is a balance
     sheet (a model check that compares exactly those rows, or a corkscrew whose closing row equals a line one of the
     totals adds); a corkscrew whose closing equals a balance sheet line is that line's roll-forward (cash -> the cash
     flow, debt, fixed assets, retained earnings); the chain of subtotals that feeds the retained earnings corkscrew is
     the P&L; the DCF's cash-flow row equals a financing row and an equity movement. When several rows satisfy an
     identity, ties break by: named by a model check, number of corkscrews tied, label vocabulary, formula shape, then
     sheet and row order. A second balance sheet or corkscrew of a kind is listed (`blocks[].alternates`), not bound.
  4. each identity is `holds`, `fails` or `unbound` (a role could not be found; the reason says which and why), and has
     a `kind`, decided per result from how its roles were bound (`bound_by` on each role: check | value_search | label |
     structure | link): a TEST when it can genuinely fail (its rows were named by a model check or by labels), or
     STRUCTURAL when it holds by construction (one row's own formula is the identity, as for the P&L subtotal chain and
     a corkscrew's closing row; or a row was chosen because its values equal its counterpart's). text() prints the tests
     first, then the structure confirmed; a structural identity still fails when a cache is stale. A model check that
     holds while our binding fails rejects our binding (the check's rows are tried; if they fail too, it is a finding);
     one that fails is reported as the model's own failure. A balance sheet whose equity is a plug (assets less
     liabilities) is unbound and a finding: it cannot fail.
  5. lineage tests follow the live edges; a stale-cache test recomputes simple arithmetic cells from the saved values
     they read, upstream of every bound row (capped, outside circular references, above a tolerance relative to the
     operands): a cell whose inputs changed after Excel last calculated will not reproduce its saved value.

Where labels decide, honestly. They are not only tie-breakers. (a) A model check is mapped to an identity by the label
vocabulary of the rows it compares (_map_given), because a check's formula alone does not say which statement it
guards. (b) The balance sheet vocabulary breaks ties between corroborated triples and names the rows when no triple
balances (the labels-only fallback binds, then TESTS: a model that does not balance is `fails`, not `unbound`).
(c) A corkscrew is called cash / debt / fixed assets / retained earnings by the words in its rows when no balance sheet
line equals it. (d) Operating / investing / financing flows, capex and distributions inside them, sources and uses, total
capex, additions and depreciation, and the P&L chain's names when no corkscrew anchors it are named by label (then by
order). Those bindings are `bound_by: label`, and an identity over them is a test. Why: a saved value cannot say which
of two equal rows is "the" distributions row; the label is the only evidence left, so it is used last and labelled.

No model (LLM) calls. The result is plain JSON values.

    uv run python -m modelatlas.statements out/<dir>/model.db [--json path] [--tolerance 1e-6]
"""
import argparse
import itertools
import json
import math
import re
import sys
import time
from collections import defaultdict
from datetime import date
from pathlib import Path

import numpy as np

from . import NOTICE, __version__, version_line
from . import dcf
from . import depgraph
from . import ontology
from . import outputs
from . import rodb

REL_TOL = 1e-6
ABS_TOL = 1e-6
MAX_POOL = 80            # candidate rows per sheet for a value search (a note says when it caps)
MAX_EVAL = 2000          # verified candidate bindings per identity
CHECKISH = re.compile(r"check|error|integrity|differen|\btie[sd]?\b|reconcil|\bok\b|\bbalance[sd]\b|\bequals?\b", re.I)
# words that only say "check" in a check's own label ('loan balances', 'account balances' and 'equals' are line items too)
STRONG_CHECK = re.compile(r"check|error|integrity|differen|reconcil|\btie[sd]?\b", re.I)
VERDICT_OK = ("OK", "TRUE", "PASS", "BALANCED")
VERDICT_WORDS = VERDICT_OK + ("FALSE", "ERROR", "ERR", "CHECK", "FAIL", "FAILED", "NOT OK", "")

# ---- R1C1 patterns -> shapes ---------------------------------------------------------------------------------------

_PATTERN = re.compile(r"(=.*?) x(\d+) \(([A-Z]+)(\d+)(?:\.\.([A-Z]+)(\d+))?\)(?:; |$)")
_SIDE = r"R(?:\[(-?\d+)\]|(\d+))?C(?:\[(-?\d+)\]|(\d+))?"
_OPERAND = re.compile(r"(?<![\w.$])((?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!)?" + _SIDE + r"(?::" + _SIDE + r")?(?![\w(\[])")


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


class Op:
    """One operand of an R1C1 formula: sheet, absolute row span, column span as ('d', delta) or ('a', col)."""
    __slots__ = ("sheet", "r1", "r2", "c1", "c2")

    def __init__(self, sheet, r1, r2, c1, c2):
        self.sheet, self.r1, self.r2, self.c1, self.c2 = sheet, r1, r2, c1, c2

    def delta(self):
        """Column offset when the operand is one column relative to the formula's own, else None."""
        return self.c1[1] if self.c1 == self.c2 and self.c1[0] == "d" else None


def _side(g, i, row, col):
    rb, ra, cb, ca = g[i], g[i + 1], g[i + 2], g[i + 3]
    r = row + int(rb) if rb is not None else (int(ra) if ra is not None else row)
    c = ("d", int(cb)) if cb is not None else (("a", int(ca)) if ca is not None else ("d", 0))
    return r, c


def parse_pattern(text: str, sheet: str, row: int):
    """The dominant pattern of a rows.patterns text -> (operands, linear terms or None, n cells). Linear terms are
    [(sign, sheet, row, col delta or None)], for formulas that are only +/- of operands and SUM() of operands."""
    m = _PATTERN.match(text or "")
    if not m:
        return None
    f, n = m.group(1), int(m.group(2))
    ops, parts = [], []

    def sub(mm):
        sh = mm.group(1)[:-1].strip("'").replace("''", "'") if mm.group(1) else sheet
        g = mm.groups()
        r1, c1 = _side(g, 1, row, 0)
        if all(x is None for x in g[5:9]) and ":" not in mm.group(0):
            r2, c2 = r1, c1
        else:
            r2, c2 = _side(g, 5, row, 0)
        ops.append(Op(sh, min(r1, r2), max(r1, r2), c1, c2))
        parts.append(len(ops) - 1)
        return f"§{len(ops) - 1}§"
    skel = _OPERAND.sub(sub, f).replace(" ", "")
    return ops, _linear(skel, ops), n


def _linear(skel: str, ops):
    """+/- chain of operands and SUM(operands) -> [(sign, operand index, inside SUM)], else None."""
    s = skel[1:] if skel.startswith("=") else skel
    tok = re.compile(r"\u00a7(\d+)\u00a7|SUM\(|[+\-]|\)|,")
    out, pos, sign, state, in_sum, sum_sign = [], 0, 1, "start", False, 1
    while pos < len(s):
        m = tok.match(s, pos)
        if not m:
            return None
        t, pos = m.group(0), m.end()
        if t in ("+", "-"):
            if in_sum:
                return None
            if state == "after":
                sign, state = (1 if t == "+" else -1), "start"
            elif t == "-":
                sign = -sign
        elif t == "SUM(":
            if in_sum or state != "start":
                return None
            in_sum, sum_sign, state = True, sign, "arg"
        elif t == ",":
            if not in_sum or state != "after_arg":
                return None
            state = "arg"
        elif t == ")":
            if not in_sum or state != "after_arg":
                return None
            in_sum, state, sign = False, "after", 1
        else:
            k = int(m.group(1))
            if in_sum:
                if state != "arg":
                    return None
                out.append((sum_sign, k, True))
                state = "after_arg"
            else:
                if state != "start":
                    return None
                out.append((sign, k, False))
                state, sign = "after", 1
    if in_sum or state != "after":
        return None
    return out


class Shape:
    """A row's dominant formula: all operands, and (when it is a plain sum) the signed rows it adds."""
    __slots__ = ("ops", "terms", "n", "prev", "has_sum", "linear")

    def __init__(self, ops, terms, n, meta, sheet):
        self.ops, self.n = ops, n
        self.linear = terms is not None
        self.terms, self.has_sum = [], False
        if terms is not None:
            for sign, k, in_sum in terms:
                op = ops[k]
                self.has_sum |= in_sum
                if op.r1 != op.r2 and not in_sum:
                    self.linear = False
                    break
                for r in range(op.r1, op.r2 + 1):
                    if (op.sheet, r) in meta:
                        self.terms.append((sign, op.sheet, r, op.delta()))
            if not self.linear:
                self.terms = []
        self.prev = {(o.sheet, r) for o in ops if o.delta() == -1 for r in range(o.r1, min(o.r2, o.r1 + 50) + 1)}


# ---- the book: model.db read once ----------------------------------------------------------------------------------

class Book:
    def __init__(self, db_path: str):
        self.db = rodb.connect(db_path)
        self.model = depgraph.Model(self.db)
        self.meta = self.model.meta
        self.layout = self.model.layout
        self.sheet_order = [s for (s,) in self.db.execute("SELECT sheet FROM sheets")]
        pats = dict(((s, r), p) for s, r, p in self.db.execute("SELECT sheet, row, patterns FROM rows"))
        self.dates: dict[str, list[tuple[int, date]]] = {}
        for s in self.sheet_order:
            lay = self.layout.get(s, {})
            hr, a, b = lay.get("header_row"), lay.get("tl_first"), lay.get("tl_last")
            if not (hr and a and b):
                continue
            ds = []
            for col, v in self.db.execute("SELECT col, value FROM cells WHERE sheet=? AND row=? AND col BETWEEN ? AND ? "
                                          "ORDER BY col", (s, hr, a, b)):
                d = dcf._as_date(v)
                if d:
                    ds.append((col, d))
            if ds:
                self.dates[s] = ds
        self.axis = sorted({d for ds in self.dates.values() for _, d in ds})
        self.axis_ix = {d: i for i, d in enumerate(self.axis)}
        self.shapes: dict[tuple, Shape] = {}
        for k, p in pats.items():
            if self.meta[k]["n_formula"]:
                r = parse_pattern(p, k[0], k[1])
                if r:
                    self.shapes[k] = Shape(r[0], r[1], r[2], self.meta, k[0])
        self._vec, self._raw = {}, {}
        self.degenerate = set()

    # -- values
    def label(self, k):
        return self.meta.get(k, {}).get("label", "")

    def name(self, k):
        return f"{k[0]}!r{k[1]}"

    def periodicity(self, sheet):
        return self.layout.get(sheet, {}).get("periodicity")

    def vec(self, k, shift: int = 0) -> np.ndarray:
        """Saved values of a row over the global date axis: NaN outside the sheet's timeline and where a period holds
        text; blank cells inside the timeline are 0. shift=1 gives the previous period's value."""
        key = (k, shift)
        if key in self._vec:
            return self._vec[key]
        out = np.full(len(self.axis), np.nan)
        ds = self.dates.get(k[0])
        if ds:
            got = dict(self.db.execute("SELECT col, value FROM cells WHERE sheet=? AND row=? AND col BETWEEN ? AND ?",
                                       (k[0], k[1], ds[0][0], ds[-1][0])))
            vals = []
            for col, d in ds:
                v = got.get(col)
                vals.append(0.0 if v is None else (_num(v) if _num(v) is not None else float("nan")))
            if shift:
                vals = [float("nan")] * shift + vals[:-shift]
            for (col, d), v in zip(ds, vals):
                out[self.axis_ix[d]] = v
        self._vec[key] = out
        return out

    def raw(self, k) -> dict:
        """Every value cell of the row right of the label and units columns, whatever the timeline: {col: value}."""
        if k not in self._raw:
            lay = self.layout.get(k[0], {})
            first = max(lay.get("label_col") or 0, lay.get("units_col") or 0)
            self._raw[k] = {c: v for c, v in self.db.execute(
                "SELECT col, value FROM cells WHERE sheet=? AND row=? AND col>? AND value IS NOT NULL", (*k, first))}
        return self._raw[k]

    def fvals(self, k) -> list:
        """The saved results of the row's formula cells (right of the label and units columns). A check's verdict is
        what its formulas compute; a typed number beside them (an enable flag, a tolerance) is not part of it."""
        lay = self.layout.get(k[0], {})
        first = max(lay.get("label_col") or 0, lay.get("units_col") or 0)
        return [v for (v,) in self.db.execute("SELECT value FROM cells WHERE sheet=? AND row=? AND col>? AND formula IS NOT "
                                              "NULL AND value IS NOT NULL ORDER BY col", (*k, first))]

    def scale(self, k) -> float:
        v = self.vec(k)
        v = v[np.isfinite(v)]
        return float(np.max(np.abs(v))) if v.size else 0.0

    def nonzero(self, k, tol=ABS_TOL) -> bool:
        return self.scale(k) > tol

    def is_static(self, k) -> bool:
        v = self.vec(k)
        v = v[np.isfinite(v)]
        return v.size > 0 and float(np.max(v) - np.min(v)) <= ABS_TOL + REL_TOL * float(np.max(np.abs(v)))

    # -- shapes
    def shape(self, k):
        return self.shapes.get(k)

    def terms(self, k):
        s = self.shapes.get(k)
        return s.terms if s and s.linear else []

    def is_total(self, k) -> bool:
        """A row that adds same-column rows: at least two terms, or a SUM() of one or more."""
        s = self.shapes.get(k)
        if not s or not s.linear or not s.terms:
            return False
        if any(t[3] != 0 for t in s.terms):
            return False
        return len(s.terms) >= 2 or s.has_sum

    def link(self, k):
        """(sign, source key) when the row is exactly one other row in the same column."""
        s = self.shapes.get(k)
        if s and s.linear and len(s.terms) == 1 and s.terms[0][3] == 0 and not s.has_sum:
            sg, sh, r, _ = s.terms[0]
            return sg, (sh, r)
        return None

    def source(self, k):
        """Follow links up to the first row that is not a link: (key, sign, path)."""
        sign, path, seen = 1, [k], {k}
        while True:
            lk = self.link(k)
            if not lk or lk[1] in seen:
                return k, sign, path
            sign *= lk[0]
            k = lk[1]
            seen.add(k)
            path.append(k)


# ---- identity evaluation ---------------------------------------------------------------------------------------------

def _tol(scale: float, tolerance) -> float:
    rel, ab = (tolerance if isinstance(tolerance, tuple) else (tolerance, ABS_TOL))
    return max(ab, rel * scale)


def evaluate(book: Book, parts, flex_keys, tolerance, scope="per_period"):
    """parts = [(key, coef, shift)]. Returns a result dict: status holds|fails|unbound, periods_checked, max_residual,
    failing_periods, signs (the sign convention that fitted), why."""
    keys = [k for k, _, _ in parts]
    sheets = {k[0] for k in keys}
    if scope == "single":
        return _evaluate_single(book, parts, tolerance)
    if len(sheets) > 1 and len({book.periodicity(s) for s in sheets}) > 1:
        return {"status": "unbound", "why": "periodicity differs (" + ", ".join(
            f"{s}: {book.periodicity(s)}" for s in sorted(sheets)) + ")"}
    vs = [book.vec(k, sh) for k, _, sh in parts]
    mask = np.all([np.isfinite(v) for v in vs], axis=0)
    n = int(mask.sum())
    if n == 0:
        return {"status": "unbound", "why": "no period with a number in every row (sheets' timelines do not overlap)"}
    scale = max(float(np.max(np.abs(v[mask]))) for v in vs)
    if scale <= ABS_TOL:
        return {"status": "unbound", "why": "every row is zero in every period (nothing to test)"}
    tol = _tol(scale, tolerance)
    flex = [i for i, (k, _, _) in enumerate(parts) if k in flex_keys]
    best = None
    for signs in itertools.product((1, -1), repeat=len(flex)):
        coef = [c for _, c, _ in parts]
        for i, s in zip(flex, signs):
            coef[i] *= s
        resid = sum(c * v for c, v in zip(coef, vs))
        r = np.abs(resid[mask])
        mx = float(r.max())
        if best is None or mx < best[0] - 1e-15:
            best = (mx, signs, resid)
        if mx <= tol:
            break
    mx, signs, resid = best
    fail = [str(book.axis[i]) for i in np.nonzero(mask & (np.abs(np.nan_to_num(resid)) > tol))[0]]
    flipped = [parts[i][0] for i, s in zip(flex, signs) if s == -1]
    return {"status": "holds" if mx <= tol else "fails", "periods_checked": n, "max_residual": mx,
            "failing_periods": fail, "flipped": flipped, "tolerance": tol}


def _evaluate_single(book: Book, parts, tolerance):
    """A block laid out in one column (sources and uses): the value column is the first column right of the label that
    holds a number in every row of the block (the next columns may be shares or notes). A block spread over several
    timeline columns of one sheet is compared column by column."""
    rows = [(k, c) for k, c, _ in parts]
    nums = [{c: _num(v) for c, v in book.raw(k).items() if _num(v) is not None} for k, _ in rows]
    if not all(nums):
        return {"status": "unbound", "why": "a row has no numbers"}
    scale = max(abs(v) for d in nums for v in d.values())
    if scale <= ABS_TOL:
        return {"status": "unbound", "why": "every row is zero (nothing to test)"}
    tol = _tol(scale, tolerance)
    common = sorted(set.intersection(*[set(d) for d in nums]))
    sheets = {k[0] for k, _ in rows}
    tl = {c for c, _ in book.dates.get(next(iter(sheets)), [])} if len(sheets) == 1 else set()
    phased = [c for c in common if c in tl]
    if len(phased) >= 2:
        res = [sum(c * d[col] for (_, c), d in zip(rows, nums)) for col in phased]
        mx = max(abs(r) for r in res)
        return {"status": "holds" if mx <= tol else "fails", "periods_checked": len(phased), "max_residual": mx,
                "failing_periods": [f"col {c}" for c, r in zip(phased, res) if abs(r) > tol], "flipped": [],
                "tolerance": tol, "scope_used": "single"}
    col = common[0] if common else None
    vals = [d[col] if col is not None else d[min(d)] for d in nums]
    tot = sum(c * v for (_, c), v in zip(rows, vals))
    return {"status": "holds" if abs(tot) <= tol else "fails", "periods_checked": 1, "max_residual": abs(tot),
            "failing_periods": [] if abs(tot) <= tol else ([f"col {col}"] if col is not None else []), "flipped": [],
            "tolerance": tol, "scope_used": "single"}


# ---- the detector ------------------------------------------------------------------------------------------------------

class Detector:
    def __init__(self, db_path: str, tolerance=None):
        self.path = db_path
        self.t0 = time.time()
        if tolerance is None:
            self.tol = (REL_TOL, ABS_TOL)
        elif isinstance(tolerance, (tuple, list)):
            self.tol = (float(tolerance[0]), float(tolerance[1]))
        elif isinstance(tolerance, dict):
            self.tol = (float(tolerance.get("rel", REL_TOL)), float(tolerance.get("abs", ABS_TOL)))
        else:
            self.tol = (float(tolerance), ABS_TOL)
        self.b = Book(db_path)
        self.bind: dict[str, object] = {}     # role -> key, or list of (key, sign) for a list role
        self.how: dict[str, str] = {}         # role -> reason
        self.why: dict[str, str] = {}         # role -> why it is not bound
        self.block_sheet: dict[str, str] = {}
        self.corkscrews: list[dict] = []
        self.subtotals: list[dict] = []
        self.consolidations: list[dict] = []
        self.checks: list[dict] = []
        self.givens: dict[str, list[dict]] = defaultdict(list)
        self.notes: list[str] = []
        self.findings: list[dict] = []
        self.results: dict[str, dict] = {}    # identity key -> result
        self.extra_parts: dict[str, list] = {}
        self.stats = {"capped": []}
        self._cv: dict = {}
        self.shift: dict[str, int] = {}
        self.bs_lines: list = []
        self.bs_alternates: list = []         # balancing triples on other sheets (another balance sheet), not bound
        self.plugs: dict = {}                 # row -> a balancing triple rejected because the row is a plug
        self.bs_plug = None
        self._tie_cache: dict = {}
        self._reads_cache: dict = {}
        self.by: dict[str, str] = {}          # role -> how it was bound: check | value_search | label | structure | link
        self.check_keys: set = set()
        self.cash_cs = None

    # ---- helpers
    def L(self, k):
        return self.b.label(k)

    def vocab_score(self, role, k):
        return 1 if ontology.vocab_hit(role, self.L(k)) else 0

    def setrole(self, role, k, how, by):
        self.bind[role] = k
        self.how[role] = how
        self.by[role] = by
        self.why.pop(role, None)

    def unbound(self, role, why):
        if role not in self.bind:
            self.why.setdefault(role, why)

    def rows_of(self, sheet):
        return [k for k in self.b.meta if k[0] == sheet]

    def close(self, a, b, allow_neg=False):
        """Rows a and b carry the same saved numbers (or minus each other) in every period they share."""
        if a == b:
            return None
        r = evaluate(self.b, [(a, 1, 0), (b, -1, 0)], set(), self.tol)
        if r["status"] == "holds":
            return 1
        if allow_neg:
            r = evaluate(self.b, [(a, 1, 0), (b, 1, 0)], set(), self.tol)
            if r["status"] == "holds":
                return -1
        return None

    # ---- stage 0: structure
    def structure(self):
        b = self.b
        # subtotals / consolidations / links
        for k in sorted(b.shapes):
            if k in self.check_keys:
                continue
            if b.is_total(k):
                ts = b.terms(k)
                sheets = {t[1] for t in ts}
                rec = {"sheet": k[0], "row": k[1], "label": self.L(k), "terms": [
                    {"sheet": t[1], "row": t[2], "label": self.L((t[1], t[2])), "sign": t[0]} for t in ts]}
                if sheets == {k[0]}:
                    self.subtotals.append(rec)
                elif len(sheets - {k[0]}) >= 2 and len(ts) >= 2:
                    self.consolidations.append(rec)
        # corkscrews
        for k in sorted(b.shapes):
            if k in self.check_keys:
                continue
            s = b.shapes[k]
            if not s.linear or len(s.terms) < 2:
                continue
            if any(t[3] not in (0, -1) for t in s.terms):
                continue
            opening, moves, selfprev = None, [], False
            for sg, sh, r, d in s.terms:
                t = (sh, r)
                if d == -1:
                    if t == k and sg == 1:
                        selfprev = True
                    else:
                        moves = None
                        break
                elif opening is None and sg == 1 and k in (b.shapes[t].prev if t in b.shapes else ()):
                    opening = t
                else:
                    moves.append((t, sg))
            if moves is None or not moves or (opening is None and not selfprev):
                continue
            self.corkscrews.append({"closing": k, "opening": opening, "movements": moves, "selfprev": selfprev})
        # the opening may be any formula that reads the closing one column back (IF(first, input, prev) is fine); the
        # CLOSING must be a plain sum of the opening and the movements (MAX(0, ...) or a closing built another way is missed)
        self.notes.append("a corkscrew's opening row must be a plain sum term of the closing row")

    # ---- stage 0b: checks
    def mine_checks(self):
        b = self.b
        cand = []
        for k, m in b.meta.items():
            if not m["n_formula"] or k not in b.shapes:
                continue
            lab = m["label"]
            if CHECKISH.search(lab) or outputs.CHECK.search(lab):
                cand.append(k)
        # rows whose saved values are all zero / TRUE / OK, that read two or more distinct rows and are built like a
        # comparison (a difference, ABS, IF, ROUND, =) are checks whatever their label; a row of zeros that merely
        # multiplies things is not (spare lines of a revenue build are zero too)
        for k, s in b.shapes.items():
            if k in cand or not b.meta[k]["n_formula"] or len(s.ops) < 2 or not b.meta[k]["label"].strip():
                continue
            vals = b.fvals(k)
            if vals and all((_num(v) is not None and abs(_num(v)) <= ABS_TOL) or str(v).strip().upper() in ("OK", "TRUE")
                            for v in vals) and not b.meta[k]["label"].lower().startswith(("opening", "closing")):
                if b.meta[k]["n_const"] == 0 and (self._compares(k) or self._diff_shape(k)) \
                        and len({(o.sheet, r) for o in s.ops for r in range(o.r1, min(o.r2, o.r1 + 20) + 1)} - {k}) >= 2:
                    cand.append(k)
        # a vocabulary-matched row that is a plain stock/flow sum is not a check ('Closing balance' is not); a row whose
        # formulas give text other than a verdict is not one (a scenario name); a row named only by a weak word ('Loan
        # balances') that computes non-zero numbers must at least be built like a difference to be one
        self.check_keys = set()
        keep, later = [], []
        for k in sorted(cand):
            lab = b.meta[k]["label"]
            vals = b.fvals(k)
            if any(_num(v) is None and str(v).strip().upper() not in VERDICT_WORDS for v in vals):
                continue
            zeroish = bool(vals) and all((_num(v) is not None and abs(_num(v)) <= max(ABS_TOL, REL_TOL * 1e3)) or
                                         str(v).strip().upper() in VERDICT_OK for v in vals)
            comp = self._compares(k)
            strong = STRONG_CHECK.search(lab) or outputs.CHECK.search(lab)
            if strong or CHECKISH.search(lab):
                if zeroish or self._diff_shape(k) or any(str(v).strip().upper() in ("ERROR", "FALSE", "CHECK") for v in vals):
                    keep.append(k)
                elif comp and (strong or self._abs_shape(k)):
                    keep.append(k)
                elif strong:
                    later.append(k)   # e.g. 'All checks' = SUM of the checks above: one if what it reads are checks
            else:
                keep.append(k)
        grew = True
        while grew:   # to a fixed point: a SUM of SUMs of checks is a check too
            grew = False
            for k in [x for x in later if x not in keep]:
                refs = {(o.sheet, r) for o in b.shapes[k].ops for r in range(o.r1, min(o.r2, o.r1 + 50) + 1)} - {k}
                refs = {t for t in refs if t in b.meta}
                if refs and refs <= set(keep):
                    keep.append(k)
                    grew = True
        self.check_keys = set(keep)
        for k in keep:
            s = b.shapes[k]
            refs = []
            for o in s.ops:
                for r in range(o.r1, min(o.r2, o.r1 + 50) + 1):
                    t = (o.sheet, r)
                    if t in b.meta and t != k and t not in refs:
                        refs.append(t)
            data = [t for t in refs if t not in self.check_keys]
            sub = [t for t in refs if t in self.check_keys]
            verdict = self._verdict(k, data)
            rec = {"sheet": k[0], "row": k[1], "label": self.L(k), "verdict": verdict,
                   "kind": "leaf" if len(data) >= 2 else ("link" if len(refs) == 1 else "aggregate"),
                   "rows": [{"sheet": t[0], "row": t[1], "label": self.L(t)} for t in data],
                   "reads_checks": [f"{t[0]}!r{t[1]}" for t in sub]}
            self.checks.append({**rec, "_data": data, "_key": k})
        self.checks.sort(key=lambda c: (c["sheet"], c["row"]))
        for c in self.checks:
            if c["kind"] == "leaf":
                self._map_given(c)

    def _compares(self, k):
        f = self.b.db.execute("SELECT formula FROM cells WHERE sheet=? AND row=? AND formula IS NOT NULL LIMIT 1", k).fetchone()
        return bool(f and re.search(r"\bIF\(|ABS\(|ROUND\(|[<>]|(?<![<>])=.*=", f[0]))

    def _abs_shape(self, k):
        f = self.b.db.execute("SELECT formula FROM cells WHERE sheet=? AND row=? AND formula IS NOT NULL LIMIT 1", k).fetchone()
        return bool(f and re.search(r"\bABS\(", f[0], re.I))

    def _diff_shape(self, k):
        s = self.b.shapes.get(k)
        return bool(s and s.linear and len(s.terms) == 2 and sum(t[0] for t in s.terms) == 0)

    def _verdict(self, k, data):
        """The model's own verdict from the check's saved values: 'holds', 'fails', or None (nothing saved)."""
        vals = self.b.fvals(k)
        if not vals:
            return None
        scale = max([self.b.scale(t) for t in data] + [0.0])
        tol = _tol(scale, self.tol)
        f = self.b.db.execute("SELECT formula FROM cells WHERE sheet=? AND row=? AND formula IS NOT NULL LIMIT 1", k).fetchone()
        boolean = bool(f and re.search(r"(?<![<>!])=(?!=)", f[0].lstrip("=")) and not re.search(r"\bIF\(|ABS\(|SUM", f[0]))
        for v in vals:
            n = _num(v)
            if n is not None:
                if boolean:
                    if n == 0:
                        return "fails"
                elif abs(n) > tol:
                    return "fails"
            else:
                t = str(v).strip().upper()
                if t in VERDICT_OK:
                    continue
                if t:
                    return "fails"
        return "holds"

    GIVEN_PAIRS = [("bs_balance", ("assets", "liabilities", "equity")), ("bs_total_le", ("assets", "total_le")),
                   ("cf_cash_tie", ("cf_closing", "bs_cash")), ("su_balance", ("sources", "uses")),
                   ("debt_tie", ("debt_closing", "bs_debt")), ("fa_tie", ("fa_closing", "bs_fixed_assets")),
                   ("capex_cf", ("capex_total", "cf_capex")), ("re_tie", ("re_closing", "bs_retained")),
                   ("rev_pnl", ("rev_total", "revenue"))]

    def _map_given(self, c):
        """Which ontology identity does this leaf check enforce? The one whose roles best match the labels of the rows
        it compares (every ordering tried); a check that matches none stays a plain check."""
        data = c["_data"]
        best = None
        for key, roles in self.GIVEN_PAIRS:
            if len(roles) != len(data):
                continue
            for perm in itertools.permutations(data):
                sc = sum(ontology.vocab_hit(r, self.L(t)) for r, t in zip(roles, perm))
                if best is None or sc > best[0]:
                    best = (sc, key, dict(zip(roles, perm)))
        if best and best[0] >= 2:
            c["identity"] = best[1]
            self.givens[best[1]].append({"check": c["_key"], "roles": best[2], "verdict": c["verdict"]})
        else:
            c["identity"] = None

    def given(self, key):
        return self.givens.get(key, [])

    # ---- stage 1: the balance sheet (the anchor: assets - liabilities - equity = 0)
    def _order(self, k):
        return (self.b.sheet_order.index(k[0]) if k[0] in self.b.sheet_order else 99, k[1])

    def _by_terms(self, T, roles):
        """The (up to) two term rows of a total, assigned to roles by vocabulary and then by row order."""
        ts = [(t[1], t[2]) for t in self.b.terms(T)]
        ts = list(dict.fromkeys(ts))
        if len(ts) != len(roles):
            return None
        best = None
        for perm in itertools.permutations(ts):
            sc = (sum(self.vocab_score(r, t) for r, t in zip(roles, perm)),
                  all(self._order(perm[i]) < self._order(perm[i + 1]) for i in range(len(perm) - 1)))
            if best is None or sc > best[0]:
                best = (sc, perm)
        return dict(zip(roles, best[1]))

    def _bs_pool(self, sheet):
        b = self.b
        out = []
        for k in self.rows_of(sheet):
            if k in self.check_keys or not b.meta[k]["n_formula"] or k not in b.shapes:
                continue
            hit = any(self.vocab_score(r, k) for r in ("assets", "liabilities", "equity", "total_le"))
            if not (b.is_total(k) or hit):
                continue
            v = b.vec(k)
            if np.isnan(v[[b.axis_ix[d] for _, d in b.dates[sheet]]]).any() or b.scale(k) <= ABS_TOL:
                continue
            out.append((hit, k))
        out.sort(key=lambda x: (-x[0], x[1]))
        pool = [k for _, k in out]
        if len(pool) > MAX_POOL:
            self.stats["capped"].append(f"balance sheet candidates on {sheet}: {len(pool)} -> {MAX_POOL}")
            pool = pool[:MAX_POOL]
        return sorted(pool)

    def _triples(self, sheet):
        b = self.b
        P = self._bs_pool(sheet)
        if len(P) < 3:
            return []
        V = np.array([b.vec(k) for k in P])
        V = V[:, np.isfinite(V).all(0)]
        if V.shape[1] < 3:
            return []
        n = len(P)
        sc = np.abs(V).max(1)
        sols = []
        for s1, s2 in ((1, 1), (-1, 1), (1, -1), (-1, -1)):
            S = s1 * V[:, None, :] + s2 * V[None, :, :]
            for a in range(n):
                tol = np.maximum(self.tol[1], self.tol[0] * np.maximum(sc[a], np.maximum(sc[:, None], sc[None, :])))
                ok = np.abs(S - V[a][None, None, :]).max(-1) <= tol
                ok &= (sc[:, None] > tol) & (sc[None, :] > tol) & (sc[a] > tol)   # each row visible at the triple's tolerance
                ok[np.arange(n), np.arange(n)] = False
                ok[a, :] = False
                ok[:, a] = False
                for i, j in zip(*np.nonzero(ok)):
                    x = self._derived(P[a], P[i], P[j])
                    if x:
                        self._note_plug(x, P[a], P[i], P[j])
                    else:
                        sols.append((P[a], P[i], P[j], s1, s2))
                if len(sols) > MAX_EVAL:
                    self.stats["capped"].append(f"balance sheet triples on {sheet} capped at {MAX_EVAL}")
                    return sols
        return sols

    def _derived(self, *ks):
        """The row that is computed from the others (its own formula adds both), else None: `net assets = total assets -
        total liabilities` makes assets = liabilities + net assets true by construction, so it is no evidence of a
        balance sheet and testing it could never fail."""
        for x in ks:
            others = {y for y in ks if y != x}
            if others and others <= {(t[1], t[2]) for t in self.b.terms(x)}:
                return x
        return None

    def _bs_lines_of(self, *tots):
        """The lines the totals add, one level down (a subtotal on the same sheet is opened)."""
        b = self.b
        out = []
        for tot in tots:
            if not tot:
                continue
            for _, sh, r, _d in b.terms(tot):
                t = (sh, r)
                if b.is_total(t) and sh == tot[0]:
                    out += [(x[1], x[2]) for x in b.terms(t)]
                else:
                    out.append(t)
        return sorted(set(out))

    def _tied(self, C, line):
        """The corkscrew closing row C and the line carry the same numbers (either sign) in every period."""
        key = (C, line)
        if key not in self._tie_cache:
            self._tie_cache[key] = bool(self.b.nonzero(C) and self.b.nonzero(line) and self.close(C, line, allow_neg=True))
        return self._tie_cache[key]

    def _corroboration(self, a, i, j):
        """Independent evidence that the balancing triple (a = i + j) is a balance sheet: a model check that compares
        exactly these three rows, or a detected corkscrew whose closing row equals a line one of the three totals adds
        (a roll-forward of a balance sheet line). A triple of totals that merely happens to balance has neither."""
        trio = {a, i, j}
        for c in self.checks:
            if c["kind"] == "leaf" and set(c["_data"]) == trio:
                return ("check", c, [])
        lines = self._bs_lines_of(a, i, j)
        ties = []
        for cs in self.corkscrews:
            C = cs["closing"]
            if C in trio:
                continue
            for line in lines:
                if line != C and self._tied(C, line):
                    ties.append((C, line))
                    break
        return ("corkscrew", None, ties) if ties else None

    def _plug_role(self, x, rows):
        """Which balance sheet role the plug row x plays: 'equity' or 'liabilities' (or 'assets' if it is the total)."""
        a, i, j = rows
        if x == a:
            return "assets"
        other = j if x == i else i
        if self.vocab_score("equity", x):
            return "equity"
        if self.vocab_score("liabilities", x):
            return "liabilities"
        return "liabilities" if self.vocab_score("equity", other) else "equity"

    def _note_plug(self, x, a, i, j, force=False):
        """A balancing triple rejected because the row x is computed from the total and the other part (equity =
        assets - liabilities): it balances by construction, so it is no balance sheet and tests nothing. Remembered,
        and only reported when no real balance sheet is found (a memo 'net assets' row beside one is harmless)."""
        if x == a or x in self.plugs:
            return
        if not force and sum(bool(self.vocab_score(r, k)) or (r == "assets" and bool(self.vocab_score("total_le", k)))
                             for r, k in (("assets", a), ("liabilities", i), ("equity", j))) < 2 and \
                sum(bool(self.vocab_score(r, k)) for r, k in (("assets", a), ("equity", i), ("liabilities", j))) < 2:
            return
        role = self._plug_role(x, (a, i, j))
        self.plugs[x] = {"row": x, "rows": (a, i, j), "role": role}

    def bind_bs(self):
        b = self.b
        A = L = E = T = None
        how, by = {}, {}
        gb, gt = self.given("bs_balance"), self.given("bs_total_le")
        uncorroborated = False
        if gb:
            r = gb[0]["roles"]
            A, L, E = r["assets"], r["liabilities"], r["equity"]
            for x in ("assets", "liabilities", "equity"):
                how[x] = f"named by the model's check {b.name(gb[0]['check'])}"
                by[x] = "check"
        elif gt:
            A, T = gt[0]["roles"]["assets"], gt[0]["roles"]["total_le"]
            how["assets"] = how["total_le"] = f"named by the model's check {b.name(gt[0]['check'])}"
            by["assets"] = by["total_le"] = "check"
            lr = self._by_terms(T, ("liabilities", "equity"))
            if lr:
                L, E = lr["liabilities"], lr["equity"]
                how["liabilities"] = how["equity"] = (f"the two rows {b.name(T)} (total liabilities and equity) adds, "
                                                      f"which the model's check {b.name(gt[0]['check'])} compares with assets")
                by["liabilities"] = by["equity"] = "structure"
        if A is None:
            # a triple of totals with assets = liabilities + equity in every period, accepted when something independent
            # of the values says it is a balance sheet: a model check naming the three rows, or a corkscrew whose closing
            # row equals a line the totals add. Labels only break ties between triples that are both corroborated.
            cands = []
            closings = {cs["closing"] for cs in self.corkscrews}
            for sheet in b.sheet_order:
                if sheet not in b.dates:
                    continue
                for a, i, j, s1, s2 in self._triples(sheet):
                    if not all(b.is_total(k) for k in (a, i, j)):
                        continue  # a balance sheet's totals are sums
                    voc = sum(self.vocab_score(r, k) for r, k in (("assets", a), ("liabilities", i), ("equity", j)))
                    if not voc and any(k in closings for k in (a, i, j)):
                        continue  # a reconciliation of three closing balances: a roll-forward's totals
                    corr = self._corroboration(a, i, j)
                    key = (bool(corr and corr[0] == "check"), len(corr[2]) if corr else 0, voc,
                           sum(b.is_total(k) for k in (a, i, j)), s1 == 1 and s2 == 1,
                           self._order(a) < self._order(i) < self._order(j), tuple(-x for x in self._order(a)),
                           tuple(-x for x in self._order(i)), tuple(-x for x in self._order(j)))
                    cands.append((key, a, i, j, sheet, corr, voc))
            good = [c for c in cands if c[5]]
            uncorroborated = bool(cands) and not good
            if good:
                key, A, L, E, sheet, corr, voc = max(good, key=lambda c: c[0])
                n = int(np.isfinite(b.vec(A)).sum())
                if corr[0] == "check":
                    ck = corr[1]
                    self.givens["bs_balance"].append({"check": ck["_key"], "roles": {"assets": A, "liabilities": L, "equity": E},
                                                      "verdict": ck["verdict"]})
                    for x in ("assets", "liabilities", "equity"):
                        how[x] = (f"assets = liabilities + equity in all {n} periods; the model's check {b.name(ck['_key'])} "
                                  f"compares exactly these three rows")
                        by[x] = "check"
                else:
                    C, line = corr[2][0]
                    for x in ("assets", "liabilities", "equity"):
                        how[x] = (f"assets = liabilities + equity in all {n} periods (no model check names them); "
                                  f"the closing row {b.name(C)} of a corkscrew equals the line {b.name(line)} the totals add")
                        by[x] = "value_search"
                # the best triple on each other sheet: another balance sheet (a per-asset one) is listed, not bound
                best_by_sheet = {}
                for c in cands:
                    if c[4] != sheet and (c[5] or c[6] >= 2) and (c[4] not in best_by_sheet or c[0] > best_by_sheet[c[4]][0]):
                        best_by_sheet[c[4]] = c
                for c in sorted(best_by_sheet.values(), key=lambda c: c[4]):
                    self.bs_alternates.append({"sheet": c[4], "assets": c[1], "liabilities": c[2], "equity": c[3],
                                               "corroborated_by": c[5][0] if c[5] else None})
        if A is None:  # labels only, to be tested (a model that does not balance is `fails`, not `unbound`)
            best = None
            for sheet in b.sheet_order:
                rows = [k for k in self.rows_of(sheet) if b.meta[k]["n_formula"] and k not in self.check_keys]
                pick = {r: max((k for k in rows if self.vocab_score(r, k)), key=lambda k: (b.is_total(k), -k[1]), default=None)
                        for r in ("assets", "liabilities", "total_le")}
                eq_c = [k for k in rows if self.vocab_score("equity", k)]
                if pick["assets"] and pick["liabilities"]:
                    for k in eq_c:
                        x = self._derived(k, pick["assets"], pick["liabilities"])
                        if x and x != pick["assets"]:
                            self._note_plug(x, pick["assets"], pick["liabilities"], k, force=True)
                pick["equity"] = max((k for k in eq_c if not (
                    pick["assets"] and pick["liabilities"] and self._derived(k, pick["assets"], pick["liabilities"]))),
                    key=lambda k: (b.is_total(k), -k[1]), default=None)
                sc = sum(v is not None for v in pick.values())
                if pick["assets"] and sc >= 3 and (best is None or sc > best[0]):
                    best = (sc, pick, sheet)
            if best:
                pick = best[1]
                A, L, E, T = pick["assets"], pick["liabilities"], pick["equity"], pick["total_le"]
                why_label = ("the combination of totals that balances is not corroborated by a model check or a corkscrew"
                             if uncorroborated else "no combination of totals balances in the saved values")
                for x in ("assets", "liabilities", "equity", "total_le"):
                    how[x] = f"named by its label only: {why_label}"
                    by[x] = "label"
                if T and not (L and E):
                    lr = self._by_terms(T, ("liabilities", "equity"))
                    if lr:
                        L, E = lr["liabilities"], lr["equity"]
                        how["liabilities"] = how["equity"] = f"the two rows {b.name(T)} (total liabilities and equity) adds"
                        by["liabilities"] = by["equity"] = "structure"
        # a plug: equity (or liabilities) computed as assets minus the other, so the balance sheet balances by construction
        if A and L and E:
            x = self._derived(A, L, E)
            if x and x != A:
                self._note_plug(x, A, L, E, force=True)
                self.bs_plug = self.plugs[x]
        if self.bs_plug is None and self.plugs and not (A and L and E):
            self.bs_plug = self.plugs[sorted(self.plugs, key=self._order)[0]]
        if A is None:
            why = ("no sheet has three totals with assets = liabilities + equity corroborated by a corkscrew or a model "
                   "check, and no row is labelled total assets / liabilities / equity with a counterpart")
            for r in ("assets", "liabilities", "equity", "total_le"):
                self.unbound(r, why)
            return
        sheet = A[0]
        self.block_sheet["balance_sheet"] = sheet
        for r, k in (("assets", A), ("liabilities", L), ("equity", E)):
            if k:
                self.setrole(r, k, how.get(r, ""), by.get(r, "label"))
            else:
                self.unbound(r, "total liabilities and equity found but it does not add exactly two rows")
        if T is None and L and E:
            for k in self.rows_of(sheet):
                if k in self.check_keys or not b.is_total(k):
                    continue
                if {(t[1], t[2]) for t in b.terms(k)} == {L, E}:
                    T, how["total_le"], by["total_le"] = k, f"{b.name(k)} adds liabilities and equity", "structure"
                    break
            if T is None:
                for k in self.rows_of(sheet):
                    if self.vocab_score("total_le", k) and b.meta[k]["n_formula"] and k not in self.check_keys:
                        T, how["total_le"], by["total_le"] = k, "named by its label (total liabilities and equity)", "label"
                        break
        if T:
            self.setrole("total_le", T, how.get("total_le", ""), by.get("total_le", "label"))
        else:
            self.unbound("total_le", "no 'total liabilities and equity' row")
        # the balance sheet's lines: what the three totals add, one level down
        self.bs_lines = [l for l in self._bs_lines_of(A, L, E) if l[0] == sheet and l not in self.check_keys]
        if not self.bs_lines:
            self.bs_lines = sorted(k for k in self.rows_of(sheet) if b.meta[k]["n_formula"] and not b.is_total(k)
                                   and k not in self.check_keys)

    # ---- stage 2: corkscrews, tied to the balance sheet by value, and what they are
    KINDS = [("cash", ("bs_cash", "cf_closing")), ("debt", ("bs_debt", "debt_closing")),
             ("fixed_assets", ("bs_fixed_assets",)), ("retained", ("bs_retained",))]
    KIND_WORDS = {"cash": r"\bcash\b", "debt": r"debt|loan|borrowing|facilit|drawdown|repay",
                  "fixed_assets": r"fixed assets|\bppe\b|deprec|addition", "retained": r"retained|npat|distribution|dividend"}

    def classify_corkscrews(self):
        b = self.b
        for cs in self.corkscrews:
            C = cs["closing"]
            cs["tie"] = None
            if self.bs_lines and b.nonzero(C):
                for line in self.bs_lines:
                    if line != C:
                        sg = self.close(C, line, allow_neg=True)
                        if sg:
                            cs["tie"] = (line, sg)
                            break
            labels = [self.L(C)] + ([self.L(cs["opening"])] if cs["opening"] else []) + [self.L(m) for m, _ in cs["movements"]]
            sc = {}
            for kind, roles in self.KINDS:
                s = 0
                if cs["tie"] and any(self.vocab_score(r, cs["tie"][0]) for r in roles[:1]):
                    s += 2
                if any(ontology.vocab_hit(r, self.L(C)) for r in roles[1:]):
                    s += 1
                s += sum(bool(re.search(self.KIND_WORDS[kind], l, re.I)) for l in labels[:1] + labels[1:])
                sc[kind] = s
            top = max(self.KINDS, key=lambda kr: sc[kr[0]])[0]
            cs["kind"] = top if sc[top] >= 1 else None
            cs["score"] = sc.get(top, 0)

    def _pick_cs(self, kind, given_key=None):
        b = self.b
        cands = [c for c in self.corkscrews if c["kind"] == kind]
        if not cands:
            return None
        gv = {g["roles"].get(r) for g in self.given(given_key) for r in g["roles"]} if given_key else set()
        cands.sort(key=lambda c: (c["closing"] in gv or bool(c["tie"] and c["tie"][0] in gv), bool(c["tie"]), c["score"],
                                  b.scale(c["closing"])), reverse=True)
        return cands[0]

    def _bind_roll(self, prefix, cs, how):
        self.setrole(f"{prefix}_closing", cs["closing"], how, "structure")
        if cs["opening"]:
            self.setrole(f"{prefix}_opening", cs["opening"], f"reads {self.b.name(cs['closing'])} one period back", "structure")
        else:
            self.setrole(f"{prefix}_opening", cs["closing"], "the closing row one period back (no opening row)", "structure")
            self.shift[f"{prefix}_opening"] = 1
        self.setrole(f"{prefix}_movements*", list(cs["movements"]), "the rows the closing row adds to the opening", "structure")

    def bind_corkscrew_blocks(self):
        b = self.b
        for kind, prefix, gk, tie_role, block in (("debt", "debt", "debt_tie", "bs_debt", "debt"),
                                                  ("fixed_assets", "fa", "fa_tie", "bs_fixed_assets", "fixed_assets"),
                                                  ("retained", "re", "re_tie", "bs_retained", "equity")):
            cs = self._pick_cs(kind, gk)
            if not cs:
                for r in (f"{prefix}_opening", f"{prefix}_movements*", f"{prefix}_closing", tie_role):
                    self.unbound(r, f"no corkscrew (closing = opening + movements, opening = closing one period back) "
                                    f"looks like {kind.replace('_', ' ')}")
                continue
            self.block_sheet[block] = cs["closing"][0]
            if cs["tie"]:
                line, sg = cs["tie"]
                how = (f"closing {b.name(cs['closing'])} equals the balance sheet line {b.name(line)}"
                       f"{' (opposite sign)' if sg == -1 else ''} in every period")
                if any(line in g["roles"].values() for g in self.given(gk)):
                    how += "; the model's own check names the pair"
                self.setrole(tie_role, line, how, self._tie_by(cs["closing"], line))
            else:
                how = f"a corkscrew whose labels read as {kind.replace('_', ' ')}; no balance sheet line equals it"
                self.unbound(tie_role, f"no balance sheet line equals the closing balance {b.name(cs['closing'])}")
            self._bind_roll(prefix, cs, how)
        cs = self._pick_cs("cash", "cf_cash_tie")
        self.cash_cs = cs
        if cs:
            if cs["tie"]:
                line, sg = cs["tie"]
                self.setrole("bs_cash", line, f"closing cash {b.name(cs['closing'])} equals balance sheet line "
                                              f"{b.name(line)} in every period", self._tie_by(cs["closing"], line))
                self.setrole("cf_closing", cs["closing"], f"equals the balance sheet's cash {b.name(line)} in every period",
                             "structure")
            else:
                self.setrole("cf_closing", cs["closing"], "a corkscrew whose label reads as cash; no balance sheet line equals it",
                             "structure")
                self.unbound("bs_cash", f"no balance sheet line equals closing cash {b.name(cs['closing'])}")
            self._bind_cf(cs)
        else:
            for r in ("cf_opening", "cf_net", "cf_closing", "cfo", "cfi", "cff", "bs_cash"):
                self.unbound(r, "no cash corkscrew (closing cash = opening cash + net flow) found")

    def _leaves(self, k):
        """The rows a statement subtotal is made of on its own sheet (links into other sheets stop the walk)."""
        b = self.b
        seen, out, stack = set(), [], [k]
        while stack:
            x = stack.pop()
            if x in seen:
                continue
            seen.add(x)
            lk = b.link(x)
            if lk and lk[1][0] == x[0]:
                stack.append(lk[1])
            elif b.is_total(x) and all(t[1] == x[0] for t in b.terms(x)):
                stack += [(t[1], t[2]) for t in b.terms(x)]
            else:
                out.append(x)
        return sorted(out)

    def _tie_by(self, C, line):
        """'link' when the balance sheet line is literally the closing row (or the reverse), else 'value_search'."""
        b = self.b
        return "link" if b.source(line)[0] == C or b.source(C)[0] == line else "value_search"

    def _bind_cf(self, cs):
        b = self.b
        self.block_sheet["cash_flow"] = cs["closing"][0]
        if cs["opening"]:
            self.setrole("cf_opening", cs["opening"], f"reads closing cash {b.name(cs['closing'])} one period back", "structure")
        else:
            self.setrole("cf_opening", cs["closing"], "closing cash one period back", "structure")
            self.shift["cf_opening"] = 1
        mv = cs["movements"]
        flows = None
        if len(mv) == 1:
            N = mv[0][0]
            self.setrole("cf_net", N, f"the one movement in closing cash {b.name(cs['closing'])}", "structure")
            flows = [(t[1], t[2]) for t in b.terms(N)]
        elif len(mv) == 3:
            flows = [m for m, _ in mv]
            self.unbound("cf_net", "closing cash adds the three flows directly, there is no net cash flow row")
        if flows and len(flows) == 3:
            best = None
            for perm in itertools.permutations(flows):
                sc = (sum(self.vocab_score(r, k) for r, k in zip(("cfo", "cfi", "cff"), perm)),
                      all(self._order(perm[i]) < self._order(perm[i + 1]) for i in range(2)))
                if best is None or sc > best[0]:
                    best = (sc, perm)
            for r, k in zip(("cfo", "cfi", "cff"), best[1]):
                self.setrole(r, k, ("named by its label" if self.vocab_score(r, k) else "by order") +
                             f"; one of the three rows {b.name(N) if len(mv) == 1 else 'closing cash'} adds",
                             "label" if self.vocab_score(r, k) else "structure")
        else:
            for r in ("cfo", "cfi", "cff"):
                self.unbound(r, "net cash flow does not add exactly three rows")
        # capex and distributions inside investing / financing
        if "cfi" in self.bind:
            c = [k for k in self._leaves(self.bind["cfi"]) if self.vocab_score("cf_capex", k)]
            if c:
                self.setrole("cf_capex", c[0], f"a row inside investing flows {b.name(self.bind['cfi'])}, labelled as capex", "label")
            else:
                self.unbound("cf_capex", "no capex-labelled row inside cash flow from investing")
        if "cff" in self.bind:
            c = [k for k in self._leaves(self.bind["cff"]) if self.vocab_score("cf_distributions", k)]
            if c:
                self.setrole("cf_distributions", c[0], f"a row inside financing flows {b.name(self.bind['cff'])}, labelled as distributions",
                             "label")
            else:
                self.unbound("cf_distributions", "no row inside cash flow from financing is labelled as distributions")

    # ---- stage 3: the P&L, a chain of subtotals that ends in the row feeding retained earnings
    CHAIN = ("ebitda", "ebit", "pbt", "npat")
    CHAIN_LEAF = {"ebit": "depreciation", "pbt": "interest", "npat": "tax"}

    def _chain_back(self, X):
        b = self.b
        chain = [X]
        while True:
            node = chain[-1]
            if not b.is_total(node):
                break
            ts = list(dict.fromkeys((t[1], t[2]) for t in b.terms(node)))
            if any(t[0] != node[0] for t in ts):
                break
            subs = [t for t in ts if b.is_total(t) and len(b.terms(t)) >= 2 and t not in chain
                    and all(x[1] == node[0] for x in b.terms(t))]
            if len(subs) != 1 or len(ts) != 2:
                break
            chain.append(subs[0])
        return chain[::-1]

    def bind_pnl(self):
        b = self.b
        X, why = None, ""
        cs = next((c for c in self.corkscrews if c["kind"] == "retained"), None) if False else self._pick_cs("retained", "re_tie")
        if cs:
            moves = [m for m, _ in cs["movements"]]
            moves.sort(key=lambda m: (-self.vocab_score("re_npat", m), self._order(m)))
            for m in moves:
                src = b.source(m)[0]
                if src != m and b.is_total(src):
                    X, why = src, f"feeds the retained earnings corkscrew {b.name(cs['closing'])} through {b.name(m)}"
                    break
        if X is None:  # no corkscrew to anchor on: the sheet with the most P&L words, the furthest chain row on it
            best = None
            for k in b.shapes:
                if b.is_total(k) and k not in self.check_keys:
                    for rank, r in enumerate(self.CHAIN):
                        if self.vocab_score(r, k):
                            hits = sum(1 for q in b.shapes if q[0] == k[0] and b.is_total(q) and
                                       any(self.vocab_score(x, q) for x in self.CHAIN))
                            key = (hits, rank, -k[1])
                            if best is None or key > best[0]:
                                best = (key, k)
            if best:
                X, why = best[1], "chosen by its label (no retained earnings corkscrew to anchor on)"
        if X is None:
            for r in self.CHAIN + ("revenue", "opex", "depreciation", "interest", "tax"):
                self.unbound(r, "no chain of subtotals ending in a row that feeds retained earnings, and no row labelled "
                                "EBITDA / EBIT / PBT / NPAT")
            return
        chain = self._chain_back(X)
        self.block_sheet["income_statement"] = X[0]
        k = len(chain)
        roles = [None] * k
        vo = [[r for r in self.CHAIN if self.vocab_score(r, n)] for n in chain]
        used = set()
        for i, n in enumerate(chain):
            cand = [r for r in vo[i] if r not in used]
            if cand:
                roles[i] = cand[0]
                used.add(cand[0])
        if any(r is None for r in roles):
            ts0 = list(dict.fromkeys((t[1], t[2]) for t in b.terms(chain[0])))
            if k == 4 and all(r in (None, c) for r, c in zip(roles, self.CHAIN)):
                roles = list(self.CHAIN)   # four subtotals, and the labelled ones already sit where the chain puts them
            else:
                roles = [r if r else None for r in roles]
                if roles[-1] is None and "feeds" in why:
                    roles[-1] = "npat"
                if roles[0] is None and "ebitda" not in roles and len(ts0) == 2:
                    sums = sorted(float(np.nansum(b.vec(t))) for t in ts0)
                    if sums[0] <= 0 <= sums[1]:  # one positive row (revenue), one negative (costs): the first margin line
                        roles[0] = "ebitda"
        for i, (n, r) in enumerate(zip(chain, roles)):
            if r:
                self.setrole(r, n, f"{'subtotal chain of ' + str(k) + ' ending at ' + b.name(chain[-1]) if k > 1 else 'a subtotal'}; "
                                   f"{why if i == k - 1 else 'a subtotal the next one adds'}", "structure")
        for r in self.CHAIN:
            if r not in self.bind:
                self.unbound(r, f"the subtotal chain has {k} rows ({', '.join(b.name(n) for n in chain)}), too few to tell "
                                f"which is {r}")
        # leaves of each node
        prev = None
        for n, r in zip(chain, roles):
            ts = list(dict.fromkeys((t[1], t[2]) for t in b.terms(n)))
            leaf = [t for t in ts if t != prev]
            if prev is None and r == "ebitda" and len(ts) == 2:
                pick = None
                for perm in itertools.permutations(ts):
                    sc = (self.vocab_score("revenue", perm[0]) + self.vocab_score("opex", perm[1]),
                          (b.vec(perm[0])[np.isfinite(b.vec(perm[0]))].sum() >= 0) and
                          (b.vec(perm[1])[np.isfinite(b.vec(perm[1]))].sum() <= 0))
                    if pick is None or sc > pick[0]:
                        pick = (sc, perm)
                rv, ox = pick[1]
                how = "named by its label" if pick[0][0] else "by sign (revenue positive, costs negative)"
                self.setrole("revenue", rv, f"{how}; adds to give {b.name(n)}", "label" if pick[0][0] else "structure")
                self.setrole("opex", ox, f"{how}; adds to give {b.name(n)}", "label" if pick[0][0] else "structure")
            elif r in self.CHAIN_LEAF and len(leaf) == 1 and prev is not None:
                self.setrole(self.CHAIN_LEAF[r], leaf[0], f"the other row {b.name(n)} adds", "structure")
            prev = n
        for r in ("revenue", "opex", "depreciation", "interest", "tax"):
            self.unbound(r, "not among the rows of the subtotal chain")

    # ---- stage 4: sources and uses, capex, revenue, equity detail, distributions
    def bind_su(self):
        b = self.b
        g = self.given("su_balance")
        if g:
            r = g[0]["roles"]
            self.setrole("sources", r["sources"], f"named by the model's check {b.name(g[0]['check'])}", "check")
            self.setrole("uses", r["uses"], f"named by the model's check {b.name(g[0]['check'])}", "check")
        else:
            best = None
            for sheet in b.sheet_order:
                rows = [k for k in self.rows_of(sheet) if b.meta[k]["n_formula"] and k not in self.check_keys and b.is_total(k)]
                s = [k for k in rows if self.vocab_score("sources", k)]
                u = [k for k in rows if self.vocab_score("uses", k)]
                if s and u:
                    best = (s[0], u[0])
                    break
            if best:
                self.setrole("sources", best[0], "named by its label (sources); no model check names it", "label")
                self.setrole("uses", best[1], "named by its label (uses); no model check names it", "label")
        if "sources" in self.bind:
            self.block_sheet["sources_uses"] = self.bind["sources"][0]
        else:
            self.unbound("sources", "no pair of totals labelled sources and uses, and no model check comparing two totals")
            self.unbound("uses", "no pair of totals labelled sources and uses")

    def bind_capex_revenue(self):
        b = self.b
        cc = self.bind.get("cf_capex")
        if cc:
            src, sg, path = b.source(cc)
            if src != cc and b.nonzero(src):
                self.setrole("capex_total", src, f"{b.name(cc)} in cash flow from investing is a link to it"
                                                 f"{' (opposite sign)' if sg == -1 else ''}", "link")
        if "capex_total" not in self.bind:
            cand = [k for k in b.shapes if b.is_total(k) and self.vocab_score("capex_total", k) and k not in self.check_keys]
            if cand:
                self.setrole("capex_total", sorted(cand)[0], "named by its label (total capex)", "label")
            else:
                self.unbound("capex_total", "the capex row in the cash flow is not a link and no total is labelled total capex")
        if "capex_total" in self.bind:
            self.block_sheet["capex"] = self.bind["capex_total"][0]
        else:
            self.unbound("cf_capex", "no capex row")
        # additions in the fixed-assets corkscrew
        mv = self.bind.get("fa_movements*")
        if mv:
            ct = self.bind.get("capex_total")
            hit = None
            if ct:
                for m, _ in mv:
                    if self.close(m, ct, allow_neg=True):
                        hit = m
                        break
            if hit:
                self.setrole("fa_additions", hit, f"equals minus/plus total capex {b.name(ct)} in every period", "value_search")
            else:
                c = [m for m, _ in mv if self.vocab_score("fa_additions", m)]
                if c:
                    self.setrole("fa_additions", c[0], "named by its label (additions)", "label")
                else:
                    self.unbound("fa_additions", "no movement of the fixed-assets corkscrew equals total capex or reads as additions")
            c = [m for m, _ in mv if self.vocab_score("fa_depreciation", m)]
            if c:
                self.setrole("fa_depreciation", c[0], "named by its label (depreciation)", "label")
            else:
                self.unbound("fa_depreciation", "no movement of the fixed-assets corkscrew reads as depreciation")
        else:
            self.unbound("fa_additions", "no fixed-assets corkscrew")
            self.unbound("fa_depreciation", "no fixed-assets corkscrew")
        # revenue
        rv = self.bind.get("revenue")
        if not rv:
            for r in ("rev_total", "rev_segments*"):
                self.unbound(r, "no P&L revenue row to trace back from")
            return
        src, sg, path = b.source(rv)
        if not b.is_total(src):  # a chain of links ending on a row that is not a sum: the revenue is built another way
            self.unbound("rev_total", f"{b.name(rv)} traces to {b.name(src)}, which is not a sum of segments")
            self.unbound("rev_segments*", "revenue is not a sum of segments")
            return
        self.setrole("rev_total", src, (f"{b.name(rv)} is a link to it" if src != rv else "the P&L revenue row itself") +
                                       " and it adds the segments", "link" if src != rv else "structure")
        self.block_sheet["revenue"] = src[0]
        segs = [((t[1], t[2]), t[0]) for t in b.terms(src)]
        self.setrole("rev_segments*", segs, f"the rows {b.name(src)} adds", "structure")

    def bind_equity_detail(self):
        b = self.b
        mv = self.bind.get("re_movements*")
        if not mv:
            for r in ("re_npat", "re_distributions"):
                self.unbound(r, "no retained earnings corkscrew")
            return
        npat = self.bind.get("npat")
        cs_sheet = self.bind["re_closing"][0]
        hit = None
        if npat:
            for m, _ in mv:
                if self.close(m, npat):
                    hit = m
                    break
        if hit:
            self.setrole("re_npat", hit, f"equals the P&L's {b.name(npat)} in every period", "value_search")
        else:
            c = [m for m, _ in mv if self.vocab_score("re_npat", m)]
            if c:
                self.setrole("re_npat", c[0], "named by its label", "label")
            else:
                self.unbound("re_npat", "no movement of the retained earnings corkscrew equals NPAT or reads as profit")
        rd = [m for m, _ in mv if self.vocab_score("re_distributions", m)]
        if rd:
            self.setrole("re_distributions", rd[0], "a movement of the corkscrew labelled as distributions", "label")
        else:
            c = [k for k in self.rows_of(cs_sheet) if self.vocab_score("re_distributions", k) and b.meta[k]["n_formula"]
                 and k not in {m for m, _ in mv} and k != self.bind["re_closing"]]
            if c:
                self.setrole("re_distributions", c[0], f"labelled as distributions on {cs_sheet}, but NOT among the rows "
                                                       f"that make up closing retained earnings", "label")
                self.findings.append({"kind": "pattern_break", "text":
                                      f"{b.name(c[0])} ({self.L(c[0])}) is not among the rows that make up closing retained "
                                      f"earnings {b.name(self.bind['re_closing'])}: distributions never reach retained earnings"})
            else:
                self.unbound("re_distributions", "no row on the equity sheet reads as distributions")

    def bind_distributions(self):
        b = self.b
        try:
            from . import valuation
            pvs = valuation._pv_cells(b.db)
        except Exception:  # noqa: BLE001
            pvs = []
        if not pvs:
            for r in ("dcf_distributions",):
                self.unbound(r, "no DCF (a SUMPRODUCT of a cash-flow row and discount factors) in the workbook")
            return
        cfrows = []
        for _, _, _, _, (a, bb), _ in pvs:
            k = (a[0], a[1])
            if k not in cfrows:
                cfrows.append(k)
        cf = self.bind.get("cf_distributions")
        eq = self.bind.get("re_distributions")
        # tie candidates: the leaves of the financing subtotal only. An enterprise DCF's cash-flow row can equal an
        # operating row by value; that is no evidence of distributions. A distributions row already found inside
        # financing is not overwritten by a value tie either.
        fin = self._leaves(self.bind["cff"]) if "cff" in self.bind else []
        for d in cfrows:
            if not b.nonzero(d):
                continue
            tie = None
            if cf and self.close(d, cf, allow_neg=True):
                tie = cf
            elif not cf:
                for k in fin:
                    if k != d and b.nonzero(k) and b.meta[k]["n_formula"] and self.close(d, k, allow_neg=True):
                        tie = k
                        break
            if not tie and not self.vocab_score("dcf_distributions", d) or (not tie and not re.search(r"distribution|dividend|equity", self.L(d), re.I)):
                self.unbound("dcf_distributions", f"the DCF discounts {b.name(d)} ({self.L(d)}), which no row of the cash flow "
                                                  f"from financing equals and which is not labelled as distributions")
                continue
            self.setrole("dcf_distributions", d, "the cash-flow row the DCF's SUMPRODUCT discounts" +
                         (f"; equals {b.name(tie)} in the cash flow in every period" if tie else ""),
                         "value_search" if tie else "label")
            self.block_sheet["distributions"] = d[0]
            if tie and tie != cf:
                self.setrole("cf_distributions", tie, f"equals the DCF's cash-flow row {b.name(d)} in every period", "value_search")
            if "cf_distributions" not in self.bind:
                self.unbound("cf_distributions", "no row of the cash flow from financing equals the DCF's cash-flow row")
            if eq is None:
                t2 = None
                for k in self.rows_of(self.bind["re_closing"][0]) if "re_closing" in self.bind else []:
                    if b.meta[k]["n_formula"] and b.nonzero(k) and self.close(d, k, allow_neg=True):
                        t2 = k
                        break
                if t2:
                    self.setrole("re_distributions", t2, f"equals the DCF's cash-flow row {b.name(d)} in every period", "value_search")
            return
        self.unbound("dcf_distributions", "the DCF's cash-flow row is all zero")

    # ---- identities
    def role_info(self, role):
        v = self.bind.get(role)
        if v is None:
            return None
        b = self.b
        if isinstance(v, list):
            out = []
            for k, s in v:
                d = {"sheet": k[0], "row": k[1], "label": self.L(k), "sign": s, "bound_by": self.by.get(role)}
                ts = b.terms(k)
                if len(ts) >= 2 and all(t[1] != k[0] for t in ts) and len({t[1] for t in ts}) >= 2:
                    d["consolidation_of"] = [{"sheet": t[1], "row": t[2], "label": self.L((t[1], t[2]))} for t in ts]
                out.append(d)
            return out
        return {"sheet": v[0], "row": v[1], "label": self.L(v), "how": self.how.get(role, ""), "bound_by": self.by.get(role)}

    def parts_for(self, ident, bind=None):
        bind = self.bind if bind is None else bind
        parts, flex, missing = [], set(), []
        for role, c in ontology.coefficients(ident).items():
            v = bind.get(role)
            if v is None:
                missing.append(role)
            elif role.endswith("*"):
                parts += [(k, c * s, 0) for k, s in v]
            else:
                parts.append((v, c, self.shift.get(role, 0)))
                if role in ident.flex:
                    flex.add(v)
        return parts, flex, missing

    def _role_rows(self, role):
        v = self.bind.get(role)
        if v is None:
            return None
        return {k for k, _ in v} if isinstance(v, list) else {v}

    def derive_kind(self, ident):
        """('test' | 'structural', why) for an identity whose roles are all bound, from how they were bound. It holds by
        construction (structural) when one row's own formula is the identity (it adds exactly the other rows), or when a
        row was chosen because its values equal its counterpart's (bound_by value_search or link). It is a test when
        every row was named independently of the identity: by a model check, by its label, or by where it sits."""
        b = self.b
        if ident.scope == "coverage":
            return "structural", "a statement about which lines have a roll-forward, not about numbers"
        roles = [r for r in ontology.coefficients(ident) if r in self.bind]
        rows = {r: self._role_rows(r) for r in roles}
        for r in roles:
            v = self.bind[r]
            if isinstance(v, list):
                continue
            others = set().union(*[rows[x] for x in roles if x != r])
            if others and others == {(t[1], t[2]) for t in b.terms(v)}:
                return "structural", f"the formula of {b.name(v)} adds exactly the other rows: it is the identity"
        found = [r for r in roles if self.by.get(r) in ("value_search", "link")]
        if found:
            return "structural", ("chosen because their values equal their counterparts': " +
                                  ", ".join(f"{r} ({self.by[r]})" for r in found))
        return "test", "rows named by " + ", ".join(sorted({self.by.get(r, "?") for r in roles}))

    def run_identity(self, ident):
        b = self.b
        rec = {"key": ident.key, "block": ident.block, "title": ident.title, "expr": ident.expr, "scope": ident.scope,
               "status": "unbound", "periods_checked": 0, "max_residual": None, "failing_periods": [],
               "roles": {r: self.role_info(r) for r in ontology.coefficients(ident) if r in self.bind},
               "given_by": None, "model_says": None, "why": "", "kind": ident.kind, "kind_basis": "not bound"}
        if ident.scope == "coverage":
            rec["kind"], rec["kind_basis"] = self.derive_kind(ident)
            return self._coverage(ident, rec)
        if ident.key in ("bs_balance", "bs_total_le") and self.bs_plug:
            rec["why"] = self._plug_reason()
            return rec
        parts, flex, missing = self.parts_for(ident)
        if missing:
            rec["why"] = "; ".join(f"{r}: {self.why.get(r, 'not found')}" for r in missing)
            return rec
        rec["kind"], rec["kind_basis"] = self.derive_kind(ident)
        res = evaluate(b, parts, flex, self.tol, ident.scope)
        gl = self.given(ident.key)
        g = None
        if gl:
            g = gl[0]
            rec["given_by"] = b.name(g["check"])
            rec["model_says"] = g["verdict"]
            same = all(self.bind.get(r) == k for r, k in g["roles"].items())
            if not same and res["status"] == "fails" and g["verdict"] == "holds":
                alt = {**self.bind, **g["roles"]}
                p2, f2, m2 = self.parts_for(ident, alt)
                r2 = evaluate(b, p2, f2, self.tol, ident.scope)
                if r2["status"] == "holds":
                    self.bind.update(g["roles"])
                    for r in g["roles"]:
                        self.how[r] = f"named by the model's check {b.name(g['check'])} (it overruled a binding found by value)"
                        self.by[r] = "check"
                    res = r2
                    rec["roles"] = {r: self.role_info(r) for r in ontology.coefficients(ident) if r in self.bind}
                    rec["kind"], rec["kind_basis"] = self.derive_kind(ident)
                else:
                    rec["why"] = ("model's own check holds; binding rejected: the rows we bound fail the identity, and so "
                                  f"do the rows the check {b.name(g['check'])} compares")
                    self.findings.append({"kind": "check_and_binding_disagree", "text":
                                          f"{ident.title}: the model's own check {b.name(g['check'])} says it holds, but "
                                          f"neither the rows we bound ({', '.join(b.name(k) for r, k in self._flat_roles(ident))}) "
                                          f"nor the rows the check compares ({', '.join(b.name(k) for k in g['roles'].values())}) "
                                          f"satisfy it in the saved values (largest residual {r2['max_residual']:.6g}): "
                                          f"the check may be switched off, loose or typed over"})
                    return rec
            elif same and res["status"] == "fails" and g["verdict"] == "holds":
                # the rows are the ones the model's own check compares, so the binding is not in doubt: the residual is
                # real and the check hides it (a looser tolerance, a switch that turns it off, or a value typed over it)
                rec["why"] = (f"fails in {len(res['failing_periods'])} of {res['periods_checked']} periods; largest residual "
                              f"{res['max_residual']:.6g}, although the model's own check {b.name(g['check'])}, which "
                              f"compares these same rows, says it holds (a looser tolerance, a switch or a typed value)")
        rec["status"] = res["status"]
        rec["why"] = rec["why"] or res.get("why", "")
        if res["status"] != "unbound":
            rec.update(periods_checked=res["periods_checked"], max_residual=res["max_residual"],
                       failing_periods=res["failing_periods"])
            if res.get("flipped"):
                rec["convention"] = "sign flipped for: " + ", ".join(
                    next(r for r in ontology.coefficients(ident) if self.bind.get(r) == k) for k in res["flipped"])
            if res.get("scope_used") and res["scope_used"] != ident.scope:
                rec["scope"] = res["scope_used"]
        if res["status"] == "holds" and not rec["why"]:
            rec["why"] = f"holds in all {res['periods_checked']} periods checked" if ident.scope != "single" else "holds"
        elif res["status"] == "fails":
            rec["why"] = rec["why"] or f"fails in {len(res['failing_periods'])} of {res['periods_checked']} periods; " \
                                       f"largest residual {res['max_residual']:.6g}"
        return rec

    def _flat_roles(self, ident):
        out = []
        for r in ontology.coefficients(ident):
            v = self.bind.get(r)
            if isinstance(v, list):
                out += [(r, k) for k, _ in v]
            elif v is not None:
                out.append((r, v))
        return out

    PLUG_FORMULA = {"equity": "assets \u2212 liabilities", "liabilities": "assets \u2212 equity", "assets": "liabilities + equity"}

    def _plug_reason(self):
        role = self.bs_plug["role"]
        return f"{role} is a plug ({self.PLUG_FORMULA[role]}), so the identity cannot test the model"

    def _coverage(self, ident, rec):
        if not self.bs_lines or "assets" not in self.bind:
            rec["why"] = "no balance sheet bound"
            return rec
        b = self.b
        tied = {cs["tie"][0] for cs in self.corkscrews if cs.get("tie")}
        un, st = [], []
        for line in self.bs_lines:
            if line in tied:
                continue
            if b.is_static(line):
                st.append(line)
            else:
                un.append(line)
        rec["periods_checked"] = len(self.bs_lines)
        if un:
            rec["why"] = "no roll-forward for " + ", ".join(f"{b.name(k)} ({self.L(k)})" for k in un)
            rec["status"] = "unbound"
        else:
            rec["status"] = "holds"
            rec["why"] = (f"{len(tied & set(self.bs_lines))} of {len(self.bs_lines)} lines tie to a corkscrew's closing balance"
                          + (f"; {len(st)} never move ({', '.join(self.L(k) for k in st)})" if st else ""))
        rec["roles"] = {"lines": [{"sheet": k[0], "row": k[1], "label": self.L(k)} for k in self.bs_lines]}
        return rec

    # ---- lineage over live edges
    def reaches(self, a, bkey, limit=20000):
        up = self.b.model.up
        seen, stack = {a}, [a]
        while stack and len(seen) < limit:
            for c in up.get(stack.pop(), ()):
                if c == bkey:
                    return True
                if c not in seen:
                    seen.add(c)
                    stack.append(c)
        return False

    def lineage(self):
        out = []

        edges = self.b.model.has_edges

        def test(key, text, reader, source):
            r, s = self.bind.get(reader), self.bind.get(source)
            if not edges:
                out.append({"key": key, "text": text, "status": "unbound",
                            "why": "model.db has no edges table (built before edges.py): rebuild it to test lineage"})
                return
            if r is None or s is None:
                miss = [x for x, v in ((reader, r), (source, s)) if v is None]
                out.append({"key": key, "text": text, "status": "unbound",
                            "why": "no row bound for " + ", ".join(miss)})
                return
            ok = self.reaches(r, s)
            out.append({"key": key, "text": text, "status": "holds" if ok else "fails",
                        "from": self.b.name(s), "to": self.b.name(r),
                        "why": f"{self.b.name(r)} reads {self.b.name(s)}" + ("" if ok else " nowhere upstream: the live edges never connect them")})
        test("revenue_to_ebitda", "P&L revenue reaches EBITDA", "ebitda", "revenue")
        test("ebitda_to_npat", "EBITDA reaches NPAT", "npat", "ebitda")
        test("npat_to_retained", "NPAT reaches closing retained earnings", "re_closing", "npat")
        test("cash_to_bs", "Closing cash reaches the balance sheet's cash", "bs_cash", "cf_closing")
        d, cf, eq = (self.bind.get(x) for x in ("dcf_distributions", "cf_distributions", "re_distributions"))
        if d and eq and not edges:
            out.append({"key": "dcf_to_equity", "text": "The DCF's cash-flow row traces back to the equity roll-forward",
                        "status": "unbound", "why": "model.db has no edges table (built before edges.py): rebuild it to "
                                                    "test lineage"})
        elif d and eq:
            via_eq = self.reaches(d, eq)
            cf_ok = (cf is None) or self.reaches(cf, eq) or self.reaches(d, cf)
            ok = via_eq or (cf and self.reaches(d, cf) and self.reaches(cf, eq))
            out.append({"key": "dcf_to_equity", "text": "The DCF's cash-flow row traces back to the equity roll-forward",
                        "status": "holds" if ok else "fails", "from": self.b.name(eq), "to": self.b.name(d),
                        "why": (f"{self.b.name(d)} reads {self.b.name(eq)}" if via_eq else
                                f"{self.b.name(d)} reaches {self.b.name(eq)} only through the cash flow" if ok else
                                f"{self.b.name(d)} never reaches {self.b.name(eq)}")
                        + ("; the cash flow's distributions " + ("come from the same row" if cf and self.reaches(cf, eq) else
                                                                  "do not read it") if cf else "")})
        else:
            out.append({"key": "dcf_to_equity", "text": "The DCF's cash-flow row traces back to the equity roll-forward",
                        "status": "unbound", "why": "no DCF row or no equity distributions row bound"})
        return out

    # ---- cheap recomputation: does a cell reproduce its saved value from the saved values it reads?
    _FUNCS = {"SUM", "ABS", "MAX", "MIN"}

    def _cellval(self, sheet, row, col):
        key = (sheet, row, col)
        if key not in self._cv:
            r = self.b.db.execute("SELECT value FROM cells WHERE sheet=? AND row=? AND col=?", key).fetchone()
            self._cv[key] = r[0] if r else None
        return self._cv[key]

    def eval_a1(self, formula, sheet):
        from openpyxl.formula import Tokenizer
        from openpyxl.formula.tokenizer import Token
        from openpyxl.utils import column_index_from_string, get_column_letter
        try:
            toks = Tokenizer(formula).items
        except Exception:  # noqa: BLE001
            return None
        # Excel's ^ is left-associative and binds looser than a leading minus (-2^2 = 4); Python's ** is neither
        powers = sum(1 for t in toks if t.type == Token.OP_IN and t.value == "^")
        if powers > 1 or (powers and any(t.type == Token.OP_PRE for t in toks)):
            return None
        out, reads, mag = [], [], 0.0
        for t in toks:
            if t.type == Token.WSPACE:
                continue
            if t.type == Token.OPERAND:
                if t.subtype == Token.NUMBER:
                    out.append(t.value)
                    mag = max(mag, abs(float(t.value)))
                elif t.subtype == Token.RANGE:
                    m = re.fullmatch(r"(?:('(?:[^']|'')+'|[A-Za-z_][\w.]*)!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?", t.value)
                    if not m:
                        return None
                    sh = m[1][1:-1].replace("''", "'") if m[1] and m[1][0] == "'" else (m[1] or sheet)
                    c1, r1 = column_index_from_string(m[2]), int(m[3])
                    c2, r2 = (column_index_from_string(m[4]), int(m[5])) if m[4] else (c1, r1)
                    if (r2 - r1 + 1) * (c2 - c1 + 1) > 400:
                        return None
                    vals = []
                    for r in range(r1, r2 + 1):
                        for c in range(c1, c2 + 1):
                            v = self._cellval(sh, r, c)
                            if v is None:
                                v = 0.0
                            elif _num(v) is None:
                                return None
                            vals.append(float(v))
                            mag = max(mag, abs(float(v)))
                            reads.append(f"{sh}!{get_column_letter(c)}{r}")
                    out.append(repr(vals[0]) if not m[4] else repr(vals))
                else:
                    return None
            elif t.type == Token.OP_IN:
                if t.value not in "+-*/^":
                    return None
                out.append("**" if t.value == "^" else t.value)
            elif t.type == Token.OP_PRE:
                out.append(t.value)
            elif t.type == Token.OP_POST:
                out.append("/100")
            elif t.type == Token.PAREN:
                out.append(t.value)
            elif t.type == Token.FUNC:
                if t.subtype == Token.OPEN:
                    name = t.value[:-1].upper()
                    if name not in self._FUNCS:
                        return None
                    out.append(f"_{name}(")
                else:
                    out.append(")")
            elif t.type == Token.SEP:
                out.append(",")
            else:
                return None

        def flat(a):
            r = []
            for x in a:
                r += flat(x) if isinstance(x, list) else [x]
            return r
        env = {"_SUM": lambda *a: sum(flat(a)), "_ABS": abs, "_MAX": lambda *a: max(flat(a)), "_MIN": lambda *a: min(flat(a))}
        try:
            v = eval("".join(out), {"__builtins__": {}}, env)  # noqa: S307 (only numbers, + - * / ** and four functions)
        except Exception:  # noqa: BLE001
            return None
        return (float(v), sorted(set(reads)), mag) if isinstance(v, (int, float)) and math.isfinite(v) else None

    STALE_CAP = 5000       # cells recomputed in all
    A1 = re.compile(r"(?:('(?:[^']|'')+'|[A-Za-z_][\w.]*)!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?(?![\w(])")

    def _reads(self, cell):
        """The cells the formula of `cell` = (sheet, row, col) reads (A1 references, ranges up to 100 cells)."""
        if cell not in self._reads_cache:
            from openpyxl.utils import column_index_from_string
            r = self.b.db.execute("SELECT formula FROM cells WHERE sheet=? AND row=? AND col=?", cell).fetchone()
            out = []
            for m in self.A1.finditer(re.sub(r'"[^"]*"', '""', r[0]) if r and r[0] else ""):
                sh = m[1][1:-1].replace("''", "'") if m[1] and m[1][0] == "'" else (m[1] or cell[0])
                c1, r1 = column_index_from_string(m[2]), int(m[3])
                c2, r2 = (column_index_from_string(m[4]), int(m[5])) if m[4] else (c1, r1)
                if (abs(c2 - c1) + 1) * (abs(r2 - r1) + 1) <= 100:
                    out += [(sh, rr, cc) for rr in range(min(r1, r2), max(r1, r2) + 1) for cc in range(min(c1, c2), max(c1, c2) + 1)]
            self._reads_cache[cell] = out
        return self._reads_cache[cell]

    def _circular(self, cell, hops=4):
        """The cell reads itself within a few hops (a circular reference: Excel's saved value is an iteration's last step,
        not what the formula gives from the saved inputs)."""
        frontier, seen = [cell], {cell}
        for _ in range(hops):
            nxt = []
            for c in frontier:
                for x in self._reads(c):
                    if x == cell:
                        return True
                    if x not in seen and len(seen) < 3000:
                        seen.add(x)
                        nxt.append(x)
            frontier = nxt
        return False

    def stale(self):
        """Does a cell reproduce its saved value from the saved values it reads? Walks upstream of every bound row, nearest
        rows first, up to STALE_CAP cells. Only a difference above max(tolerance, 1e-3 of the largest operand) is
        reported, and never on a cell in a circular reference."""
        b = self.b
        from openpyxl.utils import get_column_letter
        seeds = set()
        for v in self.bind.values():
            seeds.update([k for k, _ in v] if isinstance(v, list) else [v])
        order, seen, front = sorted(seeds), set(seeds), sorted(seeds)
        while front:
            nxt = sorted({c for k in front for c in b.model.up.get(k, ()) if c in b.meta} - seen)
            seen.update(nxt)
            order += nxt
            front = nxt
        checked = bad = circ = attempts = 0
        capped = None
        for n, k in enumerate(order):
            if checked >= self.STALE_CAP or attempts >= 4 * self.STALE_CAP:
                capped = (f"stale test stopped after {checked} cells ({len(order) - n} of {len(order)} upstream rows, the "
                          f"furthest from the bound rows, not tested)")
                break
            if not b.meta[k]["n_formula"]:
                continue
            first, skip = None, False
            for col, f, v in b.db.execute("SELECT col, formula, value FROM cells WHERE sheet=? AND row=? AND formula IS NOT NULL "
                                          "ORDER BY col LIMIT 60", k):
                saved = _num(v)
                if saved is None:
                    continue
                if checked >= self.STALE_CAP:
                    break
                attempts += 1
                r = self.eval_a1(f, k[0])
                if r is None:
                    continue
                checked += 1
                val, reads, mag = r
                if abs(val - saved) > max(self.tol[1], self.tol[0] * max(abs(val), abs(saved)), 1e-3 * mag):
                    if self._circular((k[0], k[1], col)):
                        skip = True
                        break
                    first = first or (col, f, saved, val, reads)
            if skip:
                circ += 1
            elif first:
                bad += 1
                col, f, saved, val, reads = first
                self.findings.append({"kind": "stale", "text":
                                      f"{k[0]}!{get_column_letter(col)}{k[1]} ({self.L(k)}) is saved as {saved:.6g} but the "
                                      f"values it reads, as saved, give {val:.6g} (reads {', '.join(reads[:3])}): the saved "
                                      f"results are older than an input; recalculate before trusting them"})
        self.stats["stale_cells_checked"] = checked
        self.stats["stale_rows_differing"] = bad
        self.stats["stale_rows_circular_skipped"] = circ
        if capped:
            self.stats["capped"].append(capped)

    def pattern_breaks(self):
        """A typed number in the middle of a row of formulas. A seeded row is not one: the typed first period (an opening
        balance, a base-year value) is exempt, wherever the row's first populated cell sits."""
        b = self.b
        from openpyxl.utils import get_column_letter
        seeds = set()
        for v in self.bind.values():
            seeds.update([k for k, _ in v] if isinstance(v, list) else [v])
        for k in sorted(seeds):
            ds = b.dates.get(k[0])
            if not ds or b.meta[k]["n_formula"] < 6:
                continue
            cells = list(b.db.execute("SELECT col, formula, value FROM cells WHERE sheet=? AND row=? AND col BETWEEN ? AND ? "
                                      "AND (formula IS NOT NULL OR value IS NOT NULL) ORDER BY col", (*k, ds[0][0], ds[-1][0])))
            if not cells:
                continue
            exempt = {ds[0][0], cells[0][0]}
            for col, f, v in cells:
                if f is None and col not in exempt and _num(v) not in (None, 0.0):
                    self.findings.append({"kind": "pattern_break", "text":
                                          f"{k[0]}!{get_column_letter(col)}{k[1]} ({self.L(k)}) holds the typed number {v} in a "
                                          f"row of formulas ({b.meta[k]['n_formula']} formulas)"})
                    break

    def timeline(self):
        b = self.b
        out = {}
        for s in b.sheet_order:
            ds = b.dates.get(s)
            if not ds:
                continue
            lay = b.layout[s]
            last = ds[-1][0]
            tcol = None
            h = b.db.execute("SELECT value FROM cells WHERE sheet=? AND row=? AND col=?", (s, lay["header_row"], last + 1)).fetchone()
            if h and isinstance(h[0], str) and re.search(r"total", h[0], re.I):
                tcol = last + 1
            else:
                hits = n = 0
                for r in [k for k in b.meta if k[0] == s][:200]:
                    v = b.vec(r)
                    x = b.raw(r).get(last + 1)
                    if _num(x) is None:
                        continue
                    fin = v[np.isfinite(v)]
                    if fin.size >= 3 and abs(fin.sum()) > 1e-9:
                        n += 1
                        hits += abs(fin.sum() - x) <= 1e-6 * max(1.0, abs(x))
                if n >= 3 and hits / n > 0.5:
                    tcol = last + 1
            out[s] = {"periods": len(ds), "periodicity": lay.get("periodicity"), "first": str(ds[0][1]),
                      "last": str(ds[-1][1]), "total_column": tcol}
        return out

    def cs_roll(self, cs):
        """The roll-forward of a corkscrew tested on the saved values: opening + movements - closing."""
        if "_roll" not in cs:
            parts = [(cs["opening"] or cs["closing"], 1, 0 if cs["opening"] else 1)] + [(m, s, 0) for m, s in cs["movements"]] + [(cs["closing"], -1, 0)]
            cs["_roll"] = evaluate(self.b, parts, set(), self.tol)
        return cs["_roll"]

    def cs_rec(self, cs):
        ev = self.cs_roll(cs)
        return {"closing": {"sheet": cs["closing"][0], "row": cs["closing"][1], "label": self.L(cs["closing"])},
                "opening": ({"sheet": cs["opening"][0], "row": cs["opening"][1], "label": self.L(cs["opening"])}
                            if cs["opening"] else None),
                "movements": [{"sheet": m[0], "row": m[1], "label": self.L(m), "sign": s} for m, s in cs["movements"]],
                "kind": cs["kind"], "tie": ({"sheet": cs["tie"][0][0], "row": cs["tie"][0][1],
                                            "label": self.L(cs["tie"][0])} if cs["tie"] else None),
                "status": ev["status"]}

    def alternates(self, block):
        """What else was found that could have played the block: other balance sheets, other corkscrews of its kind."""
        b = self.b
        if block == "balance_sheet":
            out = []
            for alt in self.bs_alternates:
                a, i, j = alt["assets"], alt["liabilities"], alt["equity"]
                ev = evaluate(b, [(a, 1, 0), (i, -1, 0), (j, -1, 0)], {i, j}, self.tol)
                out.append({"sheet": alt["sheet"], "corroborated_by": alt["corroborated_by"], "status": ev["status"],
                            "assets": {"sheet": a[0], "row": a[1], "label": self.L(a)},
                            "liabilities": {"sheet": i[0], "row": i[1], "label": self.L(i)},
                            "equity": {"sheet": j[0], "row": j[1], "label": self.L(j)}})
            return out
        kind, bound = {"debt": ("debt", self.bind.get("debt_closing")), "fixed_assets": ("fixed_assets", self.bind.get("fa_closing")),
                       "equity": ("retained", self.bind.get("re_closing")),
                       "cash_flow": ("cash", self.cash_cs["closing"] if self.cash_cs else None)}.get(block, (None, None))
        if kind is None:
            return []
        return [self.cs_rec(cs) for cs in sorted(self.corkscrews, key=lambda c: c["closing"])
                if cs["kind"] == kind and cs["closing"] != bound]

    # ---- everything
    def run(self):
        b = self.b
        self.mine_checks()
        self.structure()
        self.bind_bs()
        self.classify_corkscrews()
        self.bind_corkscrew_blocks()
        self.bind_pnl()
        self.bind_su()
        self.bind_capex_revenue()
        self.bind_equity_detail()
        self.bind_distributions()
        recs = [self.run_identity(i) for i in ontology.IDENTITIES]
        for r in recs:
            self.results[r["key"]] = r
        lin = self.lineage()
        self.stale()
        self.pattern_breaks()
        return recs, lin

    def result(self):
        recs, lin = self.run()
        b = self.b
        blocks = []
        for block, title in ontology.BLOCKS.items():
            names = [r.name for r in ontology.ROLES if r.block == block]
            bound = {n: self.role_info(n) for n in names if n in self.bind}
            blocks.append({"type": block, "title": title, "bound": bool(bound), "complete": len(bound) == len(names),
                           "sheet": self.block_sheet.get(block), "rows": bound,
                           "unbound": {n: self.why.get(n, "not found") for n in names if n not in self.bind},
                           "alternates": self.alternates(block)})
        cnt = {k: {"holds": 0, "fails": 0, "unbound": 0} for k in ("test", "structural")}
        for r in recs:
            cnt[r["kind"]][r["status"]] += 1
        self.stats["tests"], self.stats["structural"] = cnt["test"], cnt["structural"]
        for r in recs:
            if r["status"] == "fails" or r["model_says"] == "fails":
                self.findings.append({"kind": "model_fails_identity", "text":
                                      f"{r['title']}: {r['why']}" + (f" (the model's own check {r['given_by']} says "
                                                                       f"{r['model_says']})" if r["given_by"] else "")})
            elif r["status"] == "unbound":
                blk = next(x for x in blocks if x["type"] == r["block"])
                if blk["bound"]:
                    self.findings.append({"kind": "binding_unbound", "text": f"{r['title']}: unbound ({r['why']})"})
        for c in self.checks:
            if c["verdict"] == "fails" and c["kind"] == "leaf" and not c.get("identity"):
                self.findings.append({"kind": "model_fails_identity",
                                      "text": f"the model's check {c['sheet']}!r{c['row']} ({c['label']}) fails"})
        for x in blocks:
            if not x["bound"]:
                self.findings.append({"kind": "binding_unbound", "text": f"no {x['title'].lower()} found"})
        if self.bs_plug:
            x, role = self.bs_plug["row"], self.bs_plug["role"]
            self.findings.append({"kind": "plugged_balance_sheet", "text":
                                  f"{b.name(x)} ({self.L(x)}) is a plug: {role} is computed as {self.PLUG_FORMULA[role]}, so the "
                                  f"balance sheet balances by construction and assets = liabilities + equity cannot test the model"})
        for cs in self.corkscrews:
            ev = self.cs_roll(cs)
            if ev["status"] == "fails":
                self.findings.append({"kind": "corkscrew_fails", "text":
                                      f"{b.name(cs['closing'])} ({self.L(cs['closing'])}): opening plus movements does not give the "
                                      f"closing balance in {len(ev['failing_periods'])} of {ev['periods_checked']} periods; largest "
                                      f"residual {ev['max_residual']:.6g}"})
        order = {"model_fails_identity": 0, "plugged_balance_sheet": 1, "check_and_binding_disagree": 2, "corkscrew_fails": 3,
                 "stale": 4, "pattern_break": 5, "binding_unbound": 6}
        seen, fs = set(), []
        for f in sorted(self.findings, key=lambda f: (order.get(f["kind"], 9), f["text"])):
            if (f["kind"], f["text"]) not in seen:
                seen.add((f["kind"], f["text"]))
                fs.append(f)
        cks = [{k: v for k, v in c.items() if not k.startswith("_")} for c in self.checks]
        csr = [self.cs_rec(cs) for cs in sorted(self.corkscrews, key=lambda c: c["closing"])]
        self.stats.update({"rows": len(b.meta), "rows_parsed": len(b.shapes), "subtotals": len(self.subtotals),
                           "corkscrews": len(self.corkscrews), "consolidations": len(self.consolidations),
                           "checks": len(self.checks), "secs": round(time.time() - self.t0, 2), "notes": self.notes})
        res = {"version": __version__, "workbook": depgraph._workbook(self.path), "tolerance": {"relative": self.tol[0], "absolute": self.tol[1]},
               "timeline": self.timeline(), "blocks": blocks, "identities": recs, "corkscrews": csr,
               "subtotals": self.subtotals[:1000], "consolidations": self.consolidations[:1000], "checks": cks,
               "lineage": lin, "findings": fs, "stats": self.stats}
        return depgraph._clean(res)


def detect(db_path: str, tolerance=None) -> dict:
    """Find the statements of a model.db and test the ontology's identities on them (see the module docstring).
    tolerance: None (relative 1e-6 of the largest value, absolute floor 1e-6), a relative float, (rel, abs) or {rel, abs}."""
    return Detector(db_path, tolerance).result()


def text(result: dict) -> str:
    out = [f"{result['workbook']}: statements ({result['stats']['secs']} s), Model Atlas {result.get('version') or __version__}"]
    for blk in result["blocks"]:
        out.append(f"\n{blk['title']}" + (f"  [{blk['sheet']}]" if blk["sheet"] else "") + ("" if blk["bound"] else "  NOT FOUND"))
        for role, v in blk["rows"].items():
            if isinstance(v, list):
                out.append(f"  {role:<20} " + "; ".join(f"{x['sheet']}!r{x['row']} {x['label']}" for x in v))
            else:
                out.append(f"  {role:<20} {v['sheet']}!r{v['row']} {v['label']}  -- {v['how']}")
        for role, why in blk["unbound"].items():
            if blk["bound"]:
                out.append(f"  {role:<20} (unbound: {why})")
        for alt in blk.get("alternates", []):
            if "closing" in alt:
                out.append(f"  also found: corkscrew {alt['closing']['sheet']}!r{alt['closing']['row']} {alt['closing']['label']}"
                           f" (roll {alt['status']}{', tied to the balance sheet' if alt['tie'] else ''})")
            else:
                out.append(f"  also found: balance sheet on {alt['sheet']} (assets {alt['assets']['sheet']}!r{alt['assets']['row']}, "
                           f"{alt['status']}{', corroborated by a ' + alt['corroborated_by'] if alt['corroborated_by'] else ''})")
    for kind, head in (("test", "Tests of the model (the identity can fail: its rows were named by a check or by labels)"),
                       ("structural", "Structure confirmed (holds by construction: a row's own formula, or rows chosen "
                                      "because their values agree)")):
        out.append(f"\n{head}")
        for i in result["identities"]:
            if i["kind"] != kind:
                continue
            extra = f", max residual {i['max_residual']:.3g}" if i["max_residual"] is not None else ""
            out.append(f"  {i['status']:<8} {i['key']:<14} {i['title']}{extra}" +
                       (f" [{i['given_by']} says {i['model_says']}]" if i["given_by"] else "") +
                       (f"\n           {i['why']}" if i["status"] != "holds" and i["why"] else ""))
    out.append("\nModel's own checks")
    for c in result["checks"]:
        out.append(f"  {c['verdict'] or '?':<6} {c['sheet']}!r{c['row']} {c['label']} ({c['kind']}"
                   f"{', ' + c['identity'] if c.get('identity') else ''})")
    out.append("\nLineage")
    for l in result["lineage"]:
        out.append(f"  {l['status']:<8} {l['text']}: {l['why']}")
    out.append("\nFindings")
    for f in result["findings"]:
        out.append(f"  [{f['kind']}] {f['text']}")
    s = result["stats"]
    out.append(f"\n{s['rows']} rows, {s['rows_parsed']} formulas parsed, {s['subtotals']} subtotals, {s['corkscrews']} corkscrews, "
               f"{s['consolidations']} consolidations, {s['checks']} check rows")
    for k, name in (("tests", "tests of the model"), ("structural", "structure confirmed")):
        out.append(f"{name}: " + ", ".join(f"{n} {st}" for st, n in s[k].items()))
    if s["capped"]:
        out.append("capped: " + "; ".join(s["capped"]))
    out += ["", NOTICE]
    return "\n".join(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("db")
    ap.add_argument("--json", help="write the full result here")
    ap.add_argument("--tolerance", type=float, help="relative tolerance (default 1e-6)")
    ap.add_argument("--version", action="version", version=version_line())
    a = ap.parse_args(argv)
    res = detect(a.db, a.tolerance)
    if a.json:
        Path(a.json).write_text(json.dumps(res, indent=1, ensure_ascii=False), encoding="utf-8")
    print(text(res))


def cli() -> None:
    main()


if __name__ == "__main__":
    cli()
